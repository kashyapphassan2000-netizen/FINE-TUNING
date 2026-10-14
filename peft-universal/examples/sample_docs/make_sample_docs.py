"""Creates a small mixed document set (PDF manual, Excel parts list, Word FAQ) so the documents -> JSONL
commands in the manual run end to end. Replace examples/sample_docs/raw with YOUR documents.

python examples/sample_docs/make_sample_docs.py
"""
from pathlib import Path

import docx
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

RAW = Path(__file__).with_name("raw")
RAW.mkdir(exist_ok=True)

manual = [
    "Model P-200 centrifugal pump - Operating manual, section 4: pressure limits.",
    "The maximum continuous discharge pressure is 6 bar. Short peaks up to 7 bar are tolerated for 10 seconds.",
    "Above 7 bar the relief valve opens. If the relief valve opens more than three times per hour, stop the pump.",
    "Section 5: start-up. Prime the casing before the first start. Never run the pump dry for more than 30 seconds.",
    "Section 6: reset. To reset the controller after a fault, hold the red button for five seconds.",
    "The status LED blinks green three times when the reset succeeded. A steady red LED means the fault persists.",
    "Section 7: maintenance. Replace the mechanical seal every 4,000 operating hours.",
    "Inspect the impeller for cavitation pitting at every seal change. Pitting deeper than 1 mm requires replacement.",
]
c = canvas.Canvas(str(RAW / "P-200_manual.pdf"), pagesize=A4)
y = 800
for line in manual:
    c.drawString(40, y, line[:110]); y -= 18
    if len(line) > 110:
        c.drawString(40, y, line[110:]); y -= 18
c.save()

wb = Workbook(); ws = wb.active; ws.title = "spare_parts"
ws.append(["part_no", "description", "fits_model", "replace_every_hours", "price_eur"])
for row in [("MS-200", "mechanical seal kit", "P-200", 4000, 85), ("IMP-200", "impeller, bronze", "P-200", 20000, 240),
            ("RV-7", "relief valve 7 bar", "P-200", 15000, 60), ("GSK-5", "casing gasket", "P-200", 4000, 12),
            ("BRG-62", "motor bearing 6205", "P-200", 12000, 18)]:
    ws.append(row)
wb.save(RAW / "spare_parts.xlsx")

d = docx.Document()
d.add_heading("P-200 support FAQ", 1)
for q, a in [("Why does the relief valve open?", "Discharge pressure exceeded 7 bar; check for a closed downstream valve."),
             ("The pump is noisy at start-up. Is that normal?", "A rattling noise for the first seconds often means the casing "
              "was not primed. Stop, prime the casing, restart."),
             ("Can I run the pump dry to test it?", "No. Dry running for more than 30 seconds destroys the mechanical seal."),
             ("Who pays for a seal replaced under warranty?", "The manufacturer, if the seal failed before 4,000 hours and the "
              "service log is complete.")]:
    d.add_paragraph(q, style="List Bullet"); d.add_paragraph(a)
d.save(RAW / "support_faq.docx")
print(f"wrote 3 documents to {RAW}")

# An Excel FAQ that ALREADY has question/answer columns (kept outside raw/: it is converted directly, P06 Section 6)
faq = Workbook(); fs = faq.active; fs.title = "faq"
fs.append(["question", "answer"])
for q, a in [("What is the max pressure of the P-200?", "6 bar continuous; peaks up to 7 bar for 10 seconds."),
             ("How do I reset the controller?", "Hold the red button for five seconds."),
             ("How often is the seal replaced?", "Every 4,000 operating hours.")]:
    fs.append([q, a])
faq.save(Path(__file__).with_name("faq.xlsx"))
print(f"wrote {Path(__file__).with_name('faq.xlsx')}")
