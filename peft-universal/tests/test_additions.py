"""Tests for the manual's additions: spreadsheet/slide ingestion, the generator-prompt leak fix, full-FT and
host/time planning, records -> named SFT dataset, and the Space deployer (dry run, no network)."""
import asyncio
import json
from pathlib import Path

import pytest


def test_ingest_xlsx_tsv_pptx_keeps_short_rows(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    from peftkit.data.ingest import ingest

    src = tmp_path / "raw"
    src.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.title = "tickets"
    ws.append(["ticket", "issue", "severity"])
    for i in range(40):                                   # 40 SHORT rows: each far below min_chars
        ws.append([f"T{i}", "pump leaks at the seal" if i % 2 else "controller will not reset", i % 5 + 1])
    wb.save(src / "tickets.xlsx")
    (src / "parts.tsv").write_text("part\tprice\n" + "".join(f"seal-{i}\t{10 + i}\n" for i in range(30)))
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Maintenance schedule"
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame.text = (
        "Replace the mechanical seal every 4000 operating hours and inspect the impeller for cavitation pitting.")
    prs.save(src / "deck.pptx")

    out = tmp_path / "chunks.jsonl"
    stats = ingest(src, out, size=800, overlap=80, min_chars=100)
    rows = [json.loads(x) for x in out.read_text().splitlines()]
    text = "\n".join(r["text"] for r in rows)
    sources = {r["source"].split("#")[0] for r in rows}
    assert {"tickets.xlsx", "parts.tsv", "deck.pptx"} <= sources
    assert stats["files"] == 3
    # before the fix every 1-row text was dropped by min_chars; now all rows survive, grouped
    assert all(f"ticket: T{i} " in text for i in range(40))
    assert "[tickets] ticket: T0 | issue: controller will not reset | severity: 1" in text
    assert "seal-29" in text and "Maintenance schedule" in text


def test_build_never_writes_generator_prompt_into_training_rows(tmp_path, monkeypatch, tiny_model):
    from peftkit.data import build, synth

    seen = {}

    async def fake_call(client, model, passage, n, retries=4, gen_system=None):
        seen["gen_system"] = gen_system
        return [{"question": f"What does section {passage[:20]} say about the seal?",
                 "answer": f"It says: {passage[:60]}"}]

    monkeypatch.setattr(synth, "call_llm", fake_call)
    monkeypatch.setenv("OPENAI_BASE_URL", "http://unused/v1")
    monkeypatch.setenv("OPENAI_API_KEY", "x")
    raw = tmp_path / "raw"
    raw.mkdir()
    for i in range(12):
        (raw / f"doc{i}.txt").write_text(f"Document {i}. " + f"Seal type {i} must be replaced after {1000 + i} hours "
                                         f"of operation because wear pattern {i} appears on the face. " * 4)
    out = tmp_path / "processed"
    build.main(["--docs", str(raw), "--out", str(out), "--tokenizer", str(tiny_model), "--domain", "pump support",
                "--mode", "grounded", "--eval-ratio", "0"])
    rows = [json.loads(x) for x in (out / "train.jsonl").read_text().splitlines()]
    assert rows and "Return strict JSON" in seen["gen_system"]
    systems = {r["messages"][0]["content"] for r in rows}
    assert systems == {"You are a helpful pump support assistant."}      # NOT the generation instruction


def test_plan_full_ft_host_time_and_dataset(tiny_model, sft_files):
    from peftkit.plan import compute_flops, plan

    p = plan("configs/models/llama-3.1-8b.config.json", ["lora", "full"], 4096, 1, 16, 1e7, 1, None, None,
             "bf16", 4096, 1, gpu_prices={"H100-80": 2.0})
    lora, full = p["training"]
    assert full["trainable"] == p["counts"]["total"]
    assert full["optim_gb"] == pytest.approx(p["counts"]["total"] * 14 / 1024**3)   # 16 B/param incl. weights
    assert full["total_gb"] > 5 * lora["total_gb"]
    assert compute_flops(10, 100, full=True) == 8000 and compute_flops(10, 100) == 6000
    h100_full, h100_lora = p["time"]["full"]["H100-80"], p["time"]["lora"]["H100-80"]
    assert h100_full["gpu_hours"] > h100_lora["gpu_hours"]
    assert h100_lora["cost"] == pytest.approx(h100_lora["gpu_hours"] * 2.0, abs=0.006)  # printed hours x price
    assert "cost" not in p["time"]["lora"]["A100-80"]                          # no price given -> no fake $
    assert p["host"]["disk_gb"] > 2 * p["host"]["bf16_checkpoint_gb"]
    assert p["host_full"]["disk_gb"] > p["host"]["disk_gb"]
    d = plan(str(tiny_model), ["lora"], 512, 1, 8, None, 1, None, None, "bf16", 512, 1,
             dataset=str(sft_files / "train.jsonl"))
    assert d["dataset"]["exact"] and d["dataset"]["rows"] == 48
    assert d["compute"]["tokens_processed"] == d["dataset"]["tokens"] > 0       # counted tokens drive the time


def test_records_to_named_sft_dataset(tmp_path):
    from peftkit.data.records import build

    recs = [{"customer_tier": t, "issue_type": i, "severity": s, "resolved": s < 4, "_label_noised": True}
            for t in ("free", "pro") for i in ("billing", "outage") for s in range(1, 6)]
    src = tmp_path / "synthetic.jsonl"
    src.write_text("".join(json.dumps(r) + "\n" for r in recs))
    info = build(str(src), "tickets", "1.0", "predict", str(tmp_path / "ds"), target="resolved",
                 features=["customer_tier", "issue_type", "severity"], per_row=2)
    d = Path(info["dir"])
    rows = [json.loads(x) for f in ("train", "eval") for x in (d / f"{f}.jsonl").read_text().splitlines()]
    assert len(rows) == 40 and info["eval"] == 4
    assert all(r["messages"][-1]["content"] in ("True", "False") for r in rows)
    assert all("_label_noised" not in r["messages"][0]["content"] for r in rows)
    assert all(r["meta"]["source_type"] == "synthetic" and r["meta"]["dataset_version"] == "1.0" for r in rows)
    assert "pretty_name: tickets" in (d / "README.md").read_text()
    t = build(str(src), "tri", "1.0", "template", str(tmp_path / "ds"),
              user_t=["A {customer_tier} user has a {issue_type} issue (severity {severity}). Priority?"],
              assistant_t=["Priority {severity}."])
    first = json.loads((Path(t["dir"]) / "train.jsonl").read_text().splitlines()[0])
    assert first["messages"][-1]["content"].startswith("Priority ")


def test_space_deploy_dry_run_writes_a_valid_space():
    from peftkit.space import GRADIO_VERSION, deploy

    r = deploy("me/demo", "me/model", "endpoint", "Demo", private=True, system="You are a pump expert.",
               endpoint_url="https://api.example.com/v1", endpoint_key="secret", dry_run=True)
    folder = Path(r["folder"])
    readme = (folder / "README.md").read_text()
    assert f"sdk_version: {GRADIO_VERSION}" in readme and "app_file: app.py" in readme
    assert (folder / "app.py").exists() and (folder / "requirements.txt").read_text() == "openai>=1.0\n"
    local = Path(deploy("me/d2", "me/model", "local", "D", private=True, dry_run=True)["folder"])
    req = (local / "requirements.txt").read_text()
    assert "whl/cpu" in req and "transformers" in req          # CPU Space: CPU torch wheel, not the CUDA one
    assert r["variables"]["BACKEND"] == "endpoint" and r["secrets"] == ["ENDPOINT_API_KEY"]
    assert "secret" not in json.dumps(r["variables"])                       # keys never become plain variables


PLAIN_CHATML = ("{%- for m in messages %}{{- '<|im_start|>' + m['role'] + '\\n' + m['content'] + '<|im_end|>\\n' }}"
                "{%- endfor %}{%- if add_generation_prompt %}{{- '<|im_start|>assistant\\n' }}{%- endif %}")


def test_completion_only_loss_works_without_generation_markers(tiny_model, sft_files, tmp_path):
    """Official templates (Qwen, Llama, Mistral) have no {% generation %} markers: completion_only_loss must still
    mask the prompt, train the answer + EOS, and pass the doctor; assistant_only_loss must FAIL there."""
    import shutil

    from transformers import AutoTokenizer

    from peftkit.config import load_config
    from peftkit.diagnostics import check_masking
    from peftkit.tokenize_inspect import table
    from peftkit.train import run

    model_dir = tmp_path / "plain"
    shutil.copytree(tiny_model, model_dir)
    tok = AutoTokenizer.from_pretrained(model_dir)
    tok.chat_template = PLAIN_CHATML
    tok.save_pretrained(model_dir)
    tok = AutoTokenizer.from_pretrained(model_dir)
    from peftkit.data.templates import has_generation_markers
    assert not has_generation_markers(tok.chat_template) and has_generation_markers(PLAIN_CHATML.replace("+ m['content']", "{% generation %}{{ m['content'] }}{% endgeneration %}"))

    base = {"model": {"name_or_path": str(model_dir), "max_seq_length": 128},
            "data": {"train_file": str(sft_files / "train.jsonl"), "eval_file": str(sft_files / "eval.jsonl")},
            "train": {"output_dir": str(tmp_path / "co"), "epochs": 2, "per_device_batch_size": 8, "grad_accum": 1,
                      "lr": 5e-3, "logging_steps": 2, "eval_steps": 6, "save_steps": 1000}}
    bad = load_config("configs/base.yaml", "configs/methods/lora.yaml",
                      overrides={**base, "data": {**base["data"], "assistant_only_loss": True}})
    assert any(r.status == "FAIL" and "completion_only_loss" in r.fix for r in check_masking(bad, tok))

    cfg = load_config("configs/base.yaml", "configs/methods/lora.yaml",
                      overrides={**base, "data": {**base["data"], "completion_only_loss": True}})
    res = {r.name: r for r in check_masking(cfg, tok)}
    assert res["loss mask"].status == "PASS" and res["EOS is trained"].status == "PASS"
    assert res["prompt is a prefix"].status == "PASS"
    t = table(tok, json.loads((sft_files / "train.jsonl").read_text().splitlines()[0])["messages"], completion_only=True)
    trained = tok.decode([r["id"] for r in t["rows"] if r["trained_on"]])
    assert trained.startswith("The maximum pressure is 6 bar.") and t["eos_is_trained"]
    meta = run(cfg)
    losses = [h["loss"] for h in meta["log_history"] if "loss" in h]
    assert losses[-1] < losses[0] and meta["eval"]["eval_loss"] < 10

    with pytest.raises(ValueError):
        load_config("configs/base.yaml", overrides={"data": {"assistant_only_loss": True, "completion_only_loss": True}})


def test_full_fine_tuning_method(tiny_model, sft_files, tmp_path):
    from peftkit.config import load_config
    from peftkit.diagnostics import check_hparams, check_model
    from peftkit.generate import chat, load
    from peftkit.train import run

    cfg = load_config("configs/base.yaml", "configs/methods/full.yaml", overrides={
        "model": {"name_or_path": str(tiny_model), "max_seq_length": 128},
        "data": {"train_file": str(sft_files / "train.jsonl"), "eval_file": str(sft_files / "eval.jsonl")},
        "train": {"output_dir": str(tmp_path / "full"), "epochs": 2, "per_device_batch_size": 8, "grad_accum": 1,
                  "lr": 1e-3, "logging_steps": 2, "eval_steps": 6, "save_steps": 1000}})
    assert any(r.name == "trainable params" and r.status == "PASS" for r in check_model(cfg))
    assert {r.name: r.status for r in check_hparams(cfg)}["learning rate"] == "WARN"   # 1e-3 is not a full-FT lr
    meta = run(cfg)
    assert meta["trainable_params"] == meta["total_params"] and meta["saved_to"].endswith("model")
    losses = [h["loss"] for h in meta["log_history"] if "loss" in h]
    assert losses[-1] < losses[0]
    model, tok = load(meta["saved_to"])                     # a complete model: no merge step needed
    assert isinstance(chat(model, tok, "What is the maximum pressure?", 5), str)


@pytest.mark.parametrize("algo", ["dpo", "orpo", "kto", "grpo"])
def test_align_cli_all_algorithms(algo, tiny_model, tmp_path):
    from peftkit.align import grpo_rows, run, verifier_reward

    q = [{"role": "user", "content": "What is the maximum pressure?"}]
    good, bad = "The maximum pressure is 6 bar.", "Maybe twenty bar."
    if algo in ("dpo", "orpo"):
        rows = [{"prompt": q, "chosen": [{"role": "assistant", "content": good}],
                 "rejected": [{"role": "assistant", "content": bad}]}] * 8
    elif algo == "kto":
        rows = [{"prompt": q, "completion": [{"role": "assistant", "content": a}], "label": lab}
                for a, lab in ((good, True), (bad, False))] * 4
    else:
        rows = [{"messages": q, "expected": "6 bar", "must_include": ["6 bar"]},
                {"messages": q + [{"role": "assistant", "content": good}]}] * 4
    data = tmp_path / "d.jsonl"
    data.write_text("".join(json.dumps(r) + "\n" for r in rows))
    m = run(algo, str(tiny_model), str(data), str(tmp_path / algo), lr=1e-3, max_steps=2, batch=4 if algo != "grpo" else 4,
            grad_accum=1, r=4, max_length=128, num_generations=4, max_completion_length=8)
    assert m["global_step"] == 2 and (Path(m["adapter"]) / "adapter_config.json").exists()
    if algo == "grpo":
        g = grpo_rows(rows)
        assert g[1]["expected"] == good and g[1]["prompt"] == q                 # training rows become RL prompts
        assert verifier_reward([[{"content": "It is 6 bar."}], [{"content": "no"}]],
                               expected=["6 bar", "6 bar"], must_include=[["6 bar"], ["6 bar"]],
                               must_match=["", ""]) [0] > 0.5


def test_model_card_from_reports(tmp_path):
    from peftkit.card import build

    run = {"base_model": "Qwen/Qwen3-8B", "method": "qlora", "trainable_params": 10, "total_params": 100,
           "global_step": 5, "config": {"method": {"r": 16, "alpha": 32, "target_modules": "all-linear"},
                                        "train": {"epochs": 2, "lr": 2e-4, "per_device_batch_size": 2, "grad_accum": 8},
                                        "data": {"completion_only_loss": True}, "model": {"max_seq_length": 2048}}}
    cmp = {"rows": [{"suite": "domain", "before": 0.4, "after": 0.7, "delta": 0.3}], "verdict": "PASS: x", "n_regressions": 1}
    card = build(run, cmp, None, "me/data", "apache-2.0", "Demo", "Internal use only.")
    assert "base_model: Qwen/Qwen3-8B" in card and "| domain | 0.4 | 0.7 | +0.300 |" in card
    assert "completion-only" in card and "datasets: [me/data]" in card
    local = build({**run, "base_model": "/models/x"}, None, None, None, "apache-2.0", "Demo", "u")
    assert "base_model:" not in local and "local model" in local
