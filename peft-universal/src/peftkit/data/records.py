"""Structured records (real or synthkit-generated) -> a NAMED, versioned chat-SFT dataset.

synthkit produces rows like {"customer_tier": "pro", "severity": 4, "resolved": true}. An LLM is fine-tuned on
CHAT examples. This bridge turns every record into one or more conversations with templates you choose, stamps
lineage on every row, writes train/eval JSONL plus a dataset card, and can publish it to the Hub as a dataset.

  # an Excel/CSV FAQ that already has question/answer columns -> chat JSONL in one command:
  python -m peftkit.data.records --input faq.xlsx --name pump-faq --task template \
      --user "{question}" --assistant "{answer}" --source-type human --out data/datasets

  python -m peftkit.data.records --input reports/synth/synthetic.jsonl --name pump-tickets --version 1.0 \
      --task predict --target resolved --features customer_tier issue_type severity first_response_minutes \
      --out data/datasets

  # your own wording (Python str.format over the record's fields, any number of templates, picked per row):
  python -m peftkit.data.records --input rows.jsonl --name triage --task template \
      --user "A {customer_tier} customer reports a {issue_type} issue, severity {severity}. Route it." \
      --assistant "Route to: {queue}. Priority: {priority}." --out data/datasets

Tasks:
  predict  : "Given this record, what is <target>?" -> the target value (classification/regression as text)
  describe : the record as a JSON block -> a fluent one-paragraph description (template-based, no LLM)
  template : your --user/--assistant templates (repeat the flags to give several phrasings)
Why several phrasings: the model must learn the TASK, not one sentence. Pass 3-5 --user templates.
"""
from __future__ import annotations

import argparse
import json
import random
import time
from pathlib import Path

from .format import validate
from .lineage import stamp

PREDICT_USER = [
    "Given this record:\n{record}\nWhat is the value of {target}?",
    "{record}\n\nPredict {target} for this case.",
    "Here are the details of a case:\n{record}\nAnswer with the {target} only.",
]


def _tidy(r: dict, digits: int = 2) -> dict:
    """Round floats: 21.939982976436042 is generator noise the model would learn to imitate."""
    return {k: (round(v, digits) if isinstance(v, float) else v) for k, v in r.items()}


def _fmt_record(r: dict, fields: list[str]) -> str:
    return "\n".join(f"- {k}: {r[k]}" for k in fields if k in r and r[k] is not None)


def _describe(r: dict, fields: list[str]) -> str:
    parts = [f"{k.replace('_', ' ')} is {r[k]}" for k in fields if k in r and r[k] is not None]
    return ("This case has " + ", ".join(parts[:-1]) + (f", and {parts[-1]}" if len(parts) > 1 else parts[0]) + ".") \
        if parts else ""


def to_conversations(rows: list[dict], task: str, features: list[str] | None = None, target: str | None = None,
                     user_t: list[str] | None = None, assistant_t: list[str] | None = None,
                     system: str | None = None, per_row: int = 1, seed: int = 0,
                     label_map: dict | None = None, digits: int = 2) -> tuple[list[dict], dict]:
    rng = random.Random(seed)
    out, stats = [], {"rows": len(rows), "skipped_missing_field": 0, "invalid": 0}
    for r in rows:
        clean = _tidy({k: v for k, v in r.items() if not str(k).startswith("_")}, digits)  # drop bookkeeping cols
        feats = features or [k for k in clean if k != target]
        for _ in range(per_row):
            try:
                if task == "predict":
                    if target is None:
                        raise ValueError("--task predict needs --target")
                    if clean.get(target) is None:
                        raise KeyError(target)
                    user = rng.choice(user_t or PREDICT_USER).format(record=_fmt_record(clean, feats), target=target,
                                                                     **clean)
                    answer = (label_map or {}).get(str(clean[target]), str(clean[target]))
                elif task == "describe":
                    user = "Describe this record in one paragraph:\n" + json.dumps({k: clean[k] for k in feats if k in clean})
                    answer = _describe(clean, feats)
                elif task == "template":
                    if not (user_t and assistant_t):
                        raise ValueError("--task template needs --user and --assistant")
                    user = rng.choice(user_t).format(**clean)
                    answer = rng.choice(assistant_t).format(**clean)
                else:
                    raise ValueError(f"unknown task {task}")
            except KeyError:
                stats["skipped_missing_field"] += 1
                continue
            msgs = ([{"role": "system", "content": system}] if system else []) + [
                {"role": "user", "content": user}, {"role": "assistant", "content": answer}]
            if validate(msgs):
                stats["invalid"] += 1
                continue
            out.append({"messages": msgs})
    stats["conversations"] = len(out)
    return out, stats


CARD = """---
license: {license}
task_categories: [text-generation]
language: [en]
tags: [synthetic, fine-tuning, {name}]
pretty_name: {name}
---
# {name} (v{version})

Chat-format SFT dataset built with `peftkit.data.records` (task: `{task}`).

| item | value |
|---|---|
| source records | `{source}` ({n_rows} rows) |
| conversations | train {n_train} / eval {n_eval} |
| source type | {source_type} |
| created | {created} |
| target | {target} |
| features | {features} |

Every row carries lineage in `meta` (source, licence, version, record hash). Synthetic data: validate on REAL
held-out data before trusting any model trained on it.

Format: one JSON object per line, `{{"messages": [{{"role": "system"|"user"|"assistant", "content": "..."}}]}}`.
"""


def load_rows(path: str, sheet: str | None = None) -> list[dict]:
    """Records from .jsonl / .json / .csv / .tsv / .xlsx (first row = column names). Excel cells are read as
    their cached VALUES (formulas already computed); empty rows are skipped."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".jsonl":
        return [json.loads(x) for x in p.read_text(encoding="utf-8").splitlines() if x.strip()]
    if ext == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else [data]
    if ext in (".csv", ".tsv"):
        import csv

        with p.open(encoding="utf-8-sig", newline="") as f:
            return [dict(r) for r in csv.DictReader(f, delimiter="\t" if ext == ".tsv" else ",")]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(str(p), read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else f"col{i + 1}" for i, h in enumerate(next(rows))]
        out = [dict(zip(header, v)) for v in rows if v and any(c is not None for c in v)]
        wb.close()
        return out
    raise ValueError(f"unsupported input {ext}: use .jsonl .json .csv .tsv .xlsx")


def build(input_path: str, name: str, version: str, task: str, out_dir: str, features=None, target=None,
          user_t=None, assistant_t=None, system=None, per_row=1, eval_ratio=0.1, seed=0,
          license="other", source_type="synthetic", label_map=None, sheet=None) -> dict:
    rows = load_rows(input_path, sheet)
    convs, stats = to_conversations(rows, task, features, target, user_t, assistant_t, system, per_row, seed,
                                    label_map)
    convs = stamp(convs, source_id=f"{name}:{Path(input_path).name}", source_type=source_type, license=license,
                  dataset_version=version, transformation=f"peftkit.data.records task={task}")
    random.Random(seed).shuffle(convs)
    n_eval = int(len(convs) * eval_ratio) if len(convs) >= 20 else 0
    d = Path(out_dir) / f"{name}-v{version}"
    d.mkdir(parents=True, exist_ok=True)
    for split, data in (("eval", convs[:n_eval]), ("train", convs[n_eval:])):
        (d / f"{split}.jsonl").write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in data),
                                          encoding="utf-8")
    info = {"name": name, "version": version, "task": task, "source": input_path, **stats,
            "train": len(convs) - n_eval, "eval": n_eval, "dir": str(d)}
    (d / "dataset_info.json").write_text(json.dumps(info, indent=2))
    (d / "README.md").write_text(CARD.format(license=license, name=name, version=version, task=task,
                                             source=input_path, n_rows=len(rows), n_train=info["train"],
                                             n_eval=n_eval, source_type=source_type,
                                             created=time.strftime("%Y-%m-%d"), target=target or "-",
                                             features=", ".join(features or []) or "all fields"))
    return info


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description="records jsonl -> named chat-SFT dataset (+ optional Hub upload)")
    ap.add_argument("--input", required=True, help=".jsonl/.json/.csv/.tsv/.xlsx of records (e.g. synthkit output, an Excel FAQ)")
    ap.add_argument("--sheet", default=None, help="Excel sheet name (default: the first sheet)")
    ap.add_argument("--name", required=True, help="dataset name, e.g. pump-tickets")
    ap.add_argument("--version", default="1.0")
    ap.add_argument("--task", default="predict", choices=["predict", "describe", "template"])
    ap.add_argument("--target", default=None)
    ap.add_argument("--features", nargs="*", default=None)
    ap.add_argument("--user", action="append", default=None, help="user template; repeat for more phrasings")
    ap.add_argument("--assistant", action="append", default=None, help="assistant template; repeatable")
    ap.add_argument("--system", default=None)
    ap.add_argument("--label-map", nargs="*", default=[], metavar="VALUE=ANSWER",
                    help='e.g. True="Likely to miss the appointment" False="Likely to attend"')
    ap.add_argument("--per-row", type=int, default=1, help="conversations per record (different phrasings)")
    ap.add_argument("--eval-ratio", type=float, default=0.1)
    ap.add_argument("--license", default="other")
    ap.add_argument("--source-type", default="synthetic", choices=["synthetic", "human", "customer", "licensed", "scraped"])
    ap.add_argument("--out", default="data/datasets")
    ap.add_argument("--hub", default=None, help="also upload as a Hub DATASET repo: user/name")
    ap.add_argument("--private", action="store_true")
    a = ap.parse_args(argv)
    info = build(a.input, a.name, a.version, a.task, a.out, a.features, a.target, a.user, a.assistant, a.system,
                 a.per_row, a.eval_ratio, license=a.license, source_type=a.source_type,
                 label_map=dict(x.split("=", 1) for x in a.label_map) or None, sheet=a.sheet)
    print(json.dumps(info, indent=2))
    if a.hub:
        from ..publish import to_hub

        print("hub:", to_hub(info["dir"], a.hub, a.private, repo_type="dataset"))


if __name__ == "__main__":
    main()
